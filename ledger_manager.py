"""
Albion Online 黑市套利雷達 - 記帳本管理器
負責 記帳.xlsx 的安全讀寫與資料統計
"""

import os
import openpyxl
from datetime import datetime
from typing import List, Dict, Any

def get_ledger_file_path() -> str:
    """取得記帳檔案路徑，優先使用 data/ 目錄，若根目錄存在則沿用"""
    if os.path.exists(os.path.join("data", "記帳.xlsx")):
        return os.path.join("data", "記帳.xlsx")
    if os.path.exists("記帳.xlsx"):
        return "記帳.xlsx"
    os.makedirs("data", exist_ok=True)
    return os.path.join("data", "記帳.xlsx")

LEDGER_FILE = get_ledger_file_path()
COLUMNS = ["日期", "收支類型", "階級", "項目名稱", "品質", "數量", "單價", "稅費", "總計", "銷售方式"]

def init_ledger_file():
    """確保 記帳.xlsx 存在且有正確的標題欄位"""
    global LEDGER_FILE
    LEDGER_FILE = get_ledger_file_path()
    parent_dir = os.path.dirname(LEDGER_FILE)
    if parent_dir:
        os.makedirs(parent_dir, exist_ok=True)
    if not os.path.exists(LEDGER_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "記帳明細"
        ws.append(COLUMNS)
        wb.save(LEDGER_FILE)
    else:
        try:
            wb = openpyxl.load_workbook(LEDGER_FILE)
            ws = wb.active
            # 兼容既有 Excel：更新第 8 欄為「稅費」，若缺少第 10 欄則補上「銷售方式」
            if ws.max_column < len(COLUMNS) or ws.cell(row=1, column=8).value == "掛單費":
                ws.cell(row=1, column=8, value="稅費")
                ws.cell(row=1, column=10, value="銷售方式")
                wb.save(LEDGER_FILE)
        except Exception as e:
            print(f"Error checking ledger headers: {e}")

def get_ledger_records() -> Dict[str, Any]:
    """讀取記帳本並回傳所有記錄及收支統計"""
    init_ledger_file()
    records = []
    total_income = 0
    total_expense = 0

    try:
        wb = openpyxl.load_workbook(LEDGER_FILE, data_only=True)
        ws = wb.active
        
        # 讀取行
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) > 1:
            for idx, r in enumerate(rows[1:], start=1):
                # 忽略完全為空的行
                if not any(r):
                    continue
                
                date_val = str(r[0] if len(r) > 0 and r[0] is not None else "")
                rec_type = str(r[1] if len(r) > 1 and r[1] is not None else "支出")
                tier_val = str(r[2] if len(r) > 2 and r[2] is not None else "")
                name_val = str(r[3] if len(r) > 3 and r[3] is not None else "")
                qual_val = str(r[4] if len(r) > 4 and r[4] is not None else "一般")
                
                try:
                    qty = int(float(r[5])) if len(r) > 5 and r[5] is not None else 1
                except:
                    qty = 1
                    
                try:
                    price = int(float(r[6])) if len(r) > 6 and r[6] is not None else 0
                except:
                    price = 0
                    
                try:
                    fee = int(float(r[7])) if len(r) > 7 and r[7] is not None else 0
                except:
                    fee = 0
                    
                try:
                    total = int(float(r[8])) if len(r) > 8 and r[8] is not None else 0
                except:
                    total = (qty * price + fee) if "支出" in rec_type else (qty * price - fee)

                # 銷售方式 (掛單賣出 / 直接賣出 / 免稅 / -)
                mode_val = str(r[9]).strip() if len(r) > 9 and r[9] is not None else ""
                if not mode_val:
                    if "支出" in rec_type:
                        mode_val = "-"
                    elif fee == 0:
                        mode_val = "免稅"
                    elif (qty * price) > 0 and (fee / (qty * price)) >= 0.05:
                        mode_val = "掛單賣出"
                    else:
                        mode_val = "直接賣出"

                if "收入" in rec_type:
                    total_income += total
                else:
                    total_expense += total

                records.append({
                    "id": idx,
                    "date": date_val.split(" ")[0],
                    "type": rec_type,
                    "tier": tier_val,
                    "name": name_val,
                    "quality": qual_val,
                    "qty": qty,
                    "price": price,
                    "fee": fee,
                    "total": total,
                    "mode": mode_val
                })
    except Exception as e:
        print(f"Error reading ledger: {e}")

    net_profit = total_income - total_expense
    return {
        "records": records,
        "total_income": total_income,
        "total_expense": total_expense,
        "net_profit": net_profit
    }

def add_ledger_record(data: Dict[str, Any]) -> bool:
    """新增一筆紀錄並寫入 Excel"""
    init_ledger_file()
    try:
        wb = openpyxl.load_workbook(LEDGER_FILE)
        ws = wb.active
        
        date_str = data.get("date") or datetime.now().strftime("%Y-%m-%d")
        rec_type = data.get("type", "支出")
        tier = data.get("tier", "T4.0")
        name = data.get("name", "").strip()
        quality = data.get("quality", "一般")
        qty = int(data.get("qty", 1))
        price = int(data.get("price", 0))
        fee = int(data.get("fee", 0))
        
        raw_mode = data.get("sell_mode") or data.get("mode") or ""
        if "支出" in rec_type:
            sell_mode = "-"
        elif raw_mode in ["order", "掛單", "掛單賣出"]:
            sell_mode = "掛單賣出"
        elif raw_mode in ["instant", "直接", "直接賣出"]:
            sell_mode = "直接賣出"
        elif raw_mode in ["none", "custom", "免稅"]:
            sell_mode = "免稅"
        elif raw_mode:
            sell_mode = raw_mode
        else:
            if fee == 0:
                sell_mode = "免稅"
            elif (qty * price) > 0 and (fee / (qty * price)) >= 0.05:
                sell_mode = "掛單賣出"
            else:
                sell_mode = "直接賣出"

        if "支出" in rec_type:
            total = (qty * price) + fee
        else:
            total = (qty * price) - fee
            
        row = [date_str, rec_type, tier, name, quality, qty, price, fee, total, sell_mode]
        ws.append(row)
        wb.save(LEDGER_FILE)
        return True
    except Exception as e:
        print(f"Error adding record: {e}")
        return False

def delete_ledger_record(record_id: int) -> bool:
    """刪除指定行數 (1-based index 相對於資料行)"""
    init_ledger_file()
    try:
        wb = openpyxl.load_workbook(LEDGER_FILE)
        ws = wb.active
        
        # record_id 是第 1 筆資料，在工作表中是第 2 行 (Row 2)
        excel_row = record_id + 1
        ws.delete_rows(excel_row, 1)
        wb.save(LEDGER_FILE)
        return True
    except Exception as e:
        print(f"Error deleting record: {e}")
        return False
